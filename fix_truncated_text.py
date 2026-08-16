#!/usr/bin/env python3
"""
Fix truncated text in the presentation by re-parsing the source data file
and replacing the text content of Location and House Rules text boxes.

Also applies the two-column layout fix.
"""

import re
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.oxml.ns import qn
from pptx.text.text import _Paragraph

SRC = "/home/liang/.openclaw/workspace/coding/airbnb-ai-agent/ut_dallas_top20.txt"
PPTX = "/home/liang/.openclaw/workspace/coding/airbnb-ai-agent/ut_dallas_presentation.pptx"


def parse_listings(filepath):
    """Parse the source text file into listing dicts."""
    with open(filepath, "r") as f:
        content = f.read()

    # Split into blocks by listing number pattern at start of line
    listing_starts = []
    for m in re.finditer(r'^(\d+)\.\s+(.+?)\s+—\s+\$(\d+)/night\s*$', content, re.MULTILINE):
        listing_starts.append((m.start(), m.end(), int(m.group(1)), m.group(2), int(m.group(3))))

    listings = []
    for i, (start, end, rank, title, price) in enumerate(listing_starts):
        block_end = listing_starts[i + 1][0] if i + 1 < len(listing_starts) else len(content)
        block = content[end:block_end]

        listing = {
            "rank": rank,
            "title": title,
            "price": price,
            "url": None,
            "location": None,
            "house_rules": None,
        }

        # Extract URL
        url_m = re.search(r'URL:\s*(https?://\S+)', block)
        if url_m:
            listing["url"] = url_m.group(1)

        # Extract Location: everything from "Location:" to the next field label or end of block
        # Field labels are: "House Rules:", "Description:", at start of line with 3 spaces indent
        loc_m = re.search(r'\s+Location:\s*(.*?)(?=\n   (?:House Rules:|Description:|$))', block, re.DOTALL)
        if loc_m:
            listing["location"] = loc_m.group(1).strip()

        # Extract House Rules: everything from "House Rules:" to end of block
        # (it's always the last field in each listing)
        hr_m = re.search(r'\s+House Rules:\s*(.*)', block, re.DOTALL)
        if hr_m:
            listing["house_rules"] = hr_m.group(1).strip()

        listings.append(listing)

    return listings


def replace_text_box_body(tf, new_text):
    """Replace the body text (all paragraphs after the header) in a text frame."""
    paras = list(tf.paragraphs)
    if len(paras) < 2:
        return

    txBody = tf._txBody
    all_ps = txBody.findall(qn('a:p'))

    # Keep para 0 (header), rewrite para 1, remove rest
    body_para = paras[1]

    # Clear body para runs
    for run in list(body_para.runs):
        run._r.getparent().remove(run._r)

    lines = [l.strip() for l in new_text.split('\n') if l.strip()]

    if not lines:
        return

    # First line goes into existing body para
    run = body_para.add_run()
    run.text = lines[0]

    # Remove extra paragraphs
    for extra_p in all_ps[2:]:
        txBody.remove(extra_p)

    # Add remaining lines as new paragraphs
    for line in lines[1:]:
        new_p = txBody.makeelement(qn('a:p'), {})
        txBody.append(new_p)
        new_para = _Paragraph(new_p, txBody)
        new_run = new_para.add_run()
        new_run.text = line


def main():
    listings = parse_listings(SRC)
    print(f"Parsed {len(listings)} listings from source file")

    # Verify parsing
    for l in listings:
        loc_len = len(l["location"]) if l["location"] else 0
        hr_len = len(l["house_rules"]) if l["house_rules"] else 0
        print(f"  #{l['rank']:2d} {l['title'][:30]:30s} | loc={loc_len:4d} chars | rules={hr_len:4d} chars")

    prs = Presentation(PPTX)

    # Layout constants for two-column
    MARGIN = 0.60
    CONTENT_TOP = 2.20
    CONTENT_BOTTOM = 6.90
    CONTENT_H = CONTENT_BOTTOM - CONTENT_TOP
    COL_GAP = 0.30
    COL_W = (13.33 - 2 * MARGIN - COL_GAP) / 2
    LEFT_COL_LEFT = MARGIN
    RIGHT_COL_LEFT = MARGIN + COL_W + COL_GAP

    fixed = 0
    for idx, slide in enumerate(prs.slides):
        if idx < 2:  # Skip title and market summary slides
            continue

        listing_rank = idx - 1  # slide 3 = listing #1
        if listing_rank > len(listings):
            continue

        listing = listings[listing_rank - 1]

        location_shape = None
        rules_shape = None

        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            text = shape.text_frame.text.strip()
            if text.startswith("Location") and len(text) > len("Location"):
                location_shape = shape
            elif text.startswith("House Rules"):
                rules_shape = shape

        if location_shape and listing["location"]:
            replace_text_box_body(location_shape.text_frame, listing["location"])
            location_shape.left = Inches(LEFT_COL_LEFT)
            location_shape.top = Inches(CONTENT_TOP)
            location_shape.width = Inches(COL_W)
            location_shape.height = Inches(CONTENT_H)
            fixed += 1

        if rules_shape and listing["house_rules"]:
            replace_text_box_body(rules_shape.text_frame, listing["house_rules"])
            if location_shape:
                rules_shape.left = Inches(RIGHT_COL_LEFT)
                rules_shape.width = Inches(COL_W)
            else:
                rules_shape.left = Inches(MARGIN)
                rules_shape.width = Inches(13.33 - 2 * MARGIN)
            rules_shape.top = Inches(CONTENT_TOP)
            rules_shape.height = Inches(CONTENT_H)
            fixed += 1

    prs.save(PPTX)
    print(f"\nReplaced text on {fixed} text boxes. Saved to {PPTX}")


if __name__ == "__main__":
    main()
